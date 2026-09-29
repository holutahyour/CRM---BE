#!/usr/bin/env python3
"""
Eupepsia / Soilless Farm Lab -> Operations -> Processing reference data.

Reads the "BATCH PRODUCTION SCHEDULING" workbook and emits auditable staging JSON for
CRM.Data/Seeds/ProcessingDataSeeder:

  processing/products.json   <- PRODUCT IDENTIFICATION + PRODUCT DURATION
  processing/materials.json  <- the stock-card sheets (RAW MATERIALS x2, OTHER MATERIALS
                                INVENTORY, 2026 Raw Materials) and the produce blocks of
                                PROCESSING ACTIVITIES / 2026 Processing Activities

Reference data only: no historical rows (orders, batches, yields, stock movements) are
extracted -- the workbook stays the record of those.

EVERY value not present verbatim in the source is flagged in a `created` object so the
customer can review it (see docs/IMPORT-NOTES.md, "Processing workbook"). Names are only
merged when they differ by case/spacing or appear in the small ALIASES table below;
anything else that merely looks similar is kept separate and reported as a near-miss.

Run:  python tools/import/extract_processing.py   (needs openpyxl)
      SRC_XLSX=<path> overrides the workbook location.
"""
import os, re, json, glob, difflib
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SEED_DIR = os.path.normpath(os.path.join(HERE, "..", "..", "CRM.Data", "Seeds", "Data", "Eupepsia-Seed-Data"))
OUT = os.path.join(SEED_DIR, "processing")
DOWNLOADS = os.path.join(os.path.expanduser("~"), "Downloads")


def find_workbook():
    if os.environ.get("SRC_XLSX"):
        return os.environ["SRC_XLSX"]
    # The file name is written in Unicode bold letters; match on the plain "(1).xlsx" copy
    # whose sheets we expect rather than on the name.
    for path in glob.glob(os.path.join(DOWNLOADS, "*.xlsx")):
        try:
            wb = openpyxl.load_workbook(path, read_only=True)
        except Exception:
            continue
        if "Batch Production Schedule" in wb.sheetnames and "PRODUCT IDENTIFICATION " in wb.sheetnames:
            return path
    raise SystemExit("Batch Production Scheduling workbook not found; set SRC_XLSX.")


# Category / location names must match ProcessingCategories.cs and the seeder.
RAW = "Processing – Raw Produce"
INGREDIENTS = "Processing – Ingredients"
PACKAGING = "Processing – Packaging & Supplies"
MAIN_STORE = "Main Store"            # existing location (user decision: Warehouse/Cold Room = Main Store)
PACKAGING_STORE = "Packaging Store"  # created by the seeder

# Spelling variants merged on purpose (source spelling -> canonical). Keep this short: every
# entry is a judgement call and is reported in IMPORT-NOTES.
ALIASES = {
    "CRAY FISH": "CRAYFISH",
    "LOCUST BEAN": "LOCUST BEANS",
}

UNIT_SUFFIX = re.compile(r"\s*\((KG|G|BAGS)\)\s*$", re.IGNORECASE)
UNIT_OF = {"KG": "kg", "G": "g", "BAGS": "bags"}


def key(name):
    k = " ".join(name.upper().split())
    return ALIASES.get(k, k)


UNIT_WORDS = {"CC", "MM", "CM", "KG", "G"}


def display(name):
    """Re-case ALL-CAPS words for display ("POLYTHENE POUCH (8 x 12 inches)" -> "Polythene Pouch
    (8 x 12 inches)"); words already in mixed/lower case are kept verbatim. Returns (name, recased)."""
    name = re.sub(r"\(\s+", "(", " ".join(name.split()))

    def fix(m):
        w = m.group(0)
        if w in UNIT_WORDS:
            return w.lower()
        return w.capitalize() if len(w) > 1 else w

    fixed = re.sub(r"\b[A-Z]{2,}\b", fix, name)
    return fixed, fixed != name


def suspect_damage(raw):
    """Source titles that look like a word went missing ("Plastic  with …", "Stackable Clear s")."""
    return bool(re.search(r"\S {2,}\S", raw.strip()) or re.search(r"\s[a-z]$", raw.strip()))


def split_unit(raw):
    m = UNIT_SUFFIX.search(raw)
    if not m:
        return raw.strip(), None
    return raw[: m.start()].strip(), UNIT_OF[m.group(1).upper()]


def block_titles(ws):
    """Stock-card sheets put each material's title in a merged cell above its columns."""
    titles = set()
    for rng in ws.merged_cells.ranges:
        v = ws.cell(rng.min_row, rng.min_col).value
        if isinstance(v, str) and v.strip():
            titles.add((rng.min_row, rng.min_col, v))
    for c in ws[1]:
        if isinstance(c.value, str) and c.value.strip():
            titles.add((1, c.column, c.value))
    return [t for _, _, t in sorted(titles)]


def column_values(ws, col):
    return [str(ws.cell(r, col).value) for r in range(2, ws.max_row + 1)
            if isinstance(ws.cell(r, col).value, str) and ws.cell(r, col).value.strip()]


def main():
    src = find_workbook()
    print(f"Reading {src}")
    wb = openpyxl.load_workbook(src, data_only=True)
    os.makedirs(OUT, exist_ok=True)

    # ── Materials ──────────────────────────────────────────────────────────────
    # (sheet, titles, category, location, default unit)
    sources = [
        ("RAW MATERIALS (WAREHOUSECOLDROO", block_titles(wb["RAW MATERIALS (WAREHOUSECOLDROO"]), RAW, MAIN_STORE, "kg"),
        ("2026 Raw Materials", column_values(wb["2026 Raw Materials"], 3) + column_values(wb["2026 Raw Materials"], 10), RAW, MAIN_STORE, "kg"),
        ("PROCESSING ACTIVITIES", [t for t in block_titles(wb["PROCESSING ACTIVITIES"]) if not t.upper().startswith("YEAR")], RAW, MAIN_STORE, "kg"),
        ("2026 Processing Activities", column_values(wb["2026 Processing Activities"], 3) + column_values(wb["2026 Processing Activities"], 10), RAW, MAIN_STORE, "kg"),
        ("RAW MATERIALS (PACKAGING STORE)", block_titles(wb["RAW MATERIALS (PACKAGING STORE)"]), INGREDIENTS, PACKAGING_STORE, "kg"),
        ("OTHER MATERIALS INVENTORY", block_titles(wb["OTHER MATERIALS INVENTORY"]), PACKAGING, PACKAGING_STORE, "pcs"),
    ]

    materials = {}
    for sheet, titles, category, location, default_unit in sources:
        for raw in titles:
            base, unit = split_unit(raw)
            if not base:
                continue
            k = key(base)
            if k in materials:
                m = materials[k]
                if raw.strip() not in m["sourceNames"]:
                    m["sourceNames"].append(raw.strip())
                if sheet not in m["source"]["sheets"]:
                    m["source"]["sheets"].append(sheet)
                if unit and m["created"].get("unit_defaulted"):
                    m["unitType"], m["created"]["unit_defaulted"] = unit, False
                continue
            name, recased = display(base)
            materials[k] = {
                "name": name,
                "categoryName": category,
                "locationName": location,
                "unitType": unit or default_unit,
                "reuseSku": None,
                "sourceNames": [raw.strip()],
                "source": {"sheets": [sheet]},
                "created": {
                    "name_recased": recased,
                    "unit_defaulted": unit is None,
                    "sku_generated": True,
                    "location_assigned": location == PACKAGING_STORE and sheet == "OTHER MATERIALS INVENTORY",
                    "name_looks_damaged_in_source": suspect_damage(raw),
                },
            }

    # Aliased names: record which spelling won so the notes can say so.
    for alias, canonical in ALIASES.items():
        m = materials.get(canonical)
        if m and any(" ".join(s.upper().split()) == alias for s in m["sourceNames"]):
            m["created"]["alias_merged"] = alias

    # ── Reuse items from the earlier inventory import (exact normalized name, single match) ──
    existing = json.load(open(os.path.join(SEED_DIR, "items.json"), encoding="utf-8"))
    by_key = {}
    for it in existing:
        by_key.setdefault(key(it["name"]), []).append(it)

    near_misses = []
    existing_keys = list(by_key)
    for k, m in materials.items():
        hits = by_key.get(k, [])
        if len(hits) == 1:
            m["reuseSku"] = hits[0]["sku"]
            m["created"]["sku_generated"] = False
            m["created"]["reused_from"] = {"sku": hits[0]["sku"], "name": hits[0]["name"],
                                           "previousCategory": hits[0]["categoryName"],
                                           "quantityOnHand": hits[0]["quantityOnHand"],
                                           "unitType": hits[0]["unitType"]}
        elif len(hits) > 1:
            m["created"]["ambiguous_existing"] = [
                {"sku": h["sku"], "name": h["name"], "category": h["categoryName"]} for h in hits]
        else:
            close = difflib.get_close_matches(k, existing_keys, n=3, cutoff=0.85)
            if close:
                m["created"]["near_miss_existing"] = [by_key[c][0]["name"] for c in close]
                near_misses.append((m["name"], [by_key[c][0]["name"] for c in close]))

    # Near-misses inside the workbook itself (kept separate on purpose).
    keys = list(materials)
    for k, m in materials.items():
        close = [c for c in difflib.get_close_matches(k, keys, n=4, cutoff=0.8) if c != k]
        if close:
            m["created"]["near_miss_workbook"] = [materials[c]["name"] for c in close]

    material_list = sorted(materials.values(), key=lambda m: (m["categoryName"], m["name"].upper()))

    # ── Products ───────────────────────────────────────────────────────────────
    ident = wb["PRODUCT IDENTIFICATION "]
    duration = wb["PRODUCT DURATION"]
    durations = {}
    for r in range(2, duration.max_row + 1):
        n, d = duration.cell(r, 2).value, duration.cell(r, 3).value
        if isinstance(n, str) and n.strip() and d not in (None, ""):
            durations[key(n)] = str(d).strip()

    def cell(r, c):
        v = ident.cell(r, c).value
        return str(v).strip() if v not in (None, "") and str(v).strip() else None

    products, seen = [], set()
    for r in range(2, ident.max_row + 1):
        raw = cell(r, 2)
        if not raw:
            continue
        name, recased = display(raw)
        k = key(raw)
        seen.add(k)
        products.append({
            "name": name,
            "productCode": cell(r, 3),
            "upc": cell(r, 4),
            "sku": cell(r, 5),
            "rawMaterialId": cell(r, 6),
            "processingDuration": durations.get(k),
            "source": {"sheet": "PRODUCT IDENTIFICATION", "row": r},
            "created": {"name_recased": recased},
        })
    for k, d in durations.items():
        if k not in seen:
            name, recased = display(k)
            products.append({
                "name": name, "productCode": None, "upc": None, "sku": None, "rawMaterialId": None,
                "processingDuration": d,
                "source": {"sheet": "PRODUCT DURATION"},
                "created": {"name_recased": recased, "duration_only": True},
            })

    write("products.json", products)
    write("materials.json", material_list)

    print(f"\n{sum(1 for m in material_list if m['reuseSku'])} materials reuse an existing item; "
          f"{sum(1 for m in material_list if 'ambiguous_existing' in m['created'])} ambiguous; "
          f"{len(near_misses)} near-misses against existing items.")


def write(name, obj):
    path = os.path.join(OUT, name)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=2)
    print(f"  wrote processing/{name:20s} ({len(obj)} records)")


if __name__ == "__main__":
    main()
